import io
from pathlib import Path
from typing import List

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


class DocumentLoader:
    def __init__(self, chunk_size: int = 800, chunk_overlap: int = 150):
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", "。", "；", "，", " ", ""],
        )

    def load_from_bytes(self, content: bytes, filename: str, metadata: dict = None) -> List[Document]:
        suffix = Path(filename).suffix.lower()
        text = self._extract_text(content, suffix, filename)
        base_meta = {"source": filename, "file_type": suffix}
        if metadata:
            base_meta.update(metadata)
        doc = Document(page_content=text, metadata=base_meta)
        return self.splitter.split_documents([doc])

    def load_from_text(self, text: str, source: str = "direct_input", metadata: dict = None) -> List[Document]:
        base_meta = {"source": source, "file_type": "text"}
        if metadata:
            base_meta.update(metadata)
        doc = Document(page_content=text, metadata=base_meta)
        return self.splitter.split_documents([doc])

    def _extract_text(self, content: bytes, suffix: str, filename: str) -> str:
        if suffix == ".pdf":
            return self._extract_pdf(content)
        elif suffix in (".docx", ".doc"):
            return self._extract_docx(content)
        elif suffix in (".txt", ".md"):
            return content.decode("utf-8", errors="ignore")
        else:
            return content.decode("utf-8", errors="ignore")

    def _extract_pdf(self, content: bytes) -> str:
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            return "\n\n".join(pages)
        except Exception as e:
            return f"[PDF解析错误: {e}]"

    def _extract_docx(self, content: bytes) -> str:
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)
        except Exception as e:
            return f"[DOCX解析错误: {e}]"
